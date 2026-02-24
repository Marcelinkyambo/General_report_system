import streamlit as st
import mysql.connector
import pandas as pd
import numpy as np
import hashlib
import plotly.express as px
from io import BytesIO
from fpdf import FPDF
import openpyxl  # ensure ExcelWriter engine is available
import re
from typing import Iterable, Optional
from datetime import datetime, date  # Added for dynamic date calculations
from pathlib import Path

# ----------------- Database Connection ----------------- #
def connect_to_database():
    return mysql.connector.connect(
        host="localhost",
        user="root",
        password="Root@123",
        database="data_analysis",
    )

_ALLOWED_TABLES: set[str] = {
    "stock_summary",
    "sales_by_item",
    "purchases_by_item",
    "stock_summary_previous",
    "sales_by_item_previous",
    "purchases_by_item_previous",
}


def _normalize_col_name(name: object) -> str:
    s = str(name).strip().lower()
    s = s.replace("%", "pct")
    s = re.sub(r"[^\w]+", "_", s)
    return s.strip("_")


def _find_col(df: pd.DataFrame, must_have: Iterable[str], must_not_have: Iterable[str] = ()) -> Optional[str]:
    if df.empty:
        return None
    must_have = [_normalize_col_name(x) for x in must_have]
    must_not_have = [_normalize_col_name(x) for x in must_not_have]

    normalized = {c: _normalize_col_name(c) for c in df.columns}
    for original, norm in normalized.items():
        if all(k in norm for k in must_have) and not any(k in norm for k in must_not_have):
            return original
    return None


def _normalize_sku(sku: object) -> str:
    if sku is None or (isinstance(sku, float) and pd.isna(sku)) or pd.isna(sku):
        return ""
    s = str(sku).strip()
    # Handle numeric coercion like "111-001-20.0"
    if s.endswith(".0"):
        s = s[:-2]
    return s


def _sku_prefix(sku: str, n: int = 6) -> str:
    """Return the first *n* digit characters from *sku* for grouping."""
    digits = re.sub(r"[^0-9]", "", str(sku))
    return digits[:n] if digits else str(sku).strip()[:n]


def _to_number(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce").fillna(0)


@st.cache_data(show_spinner=False)
def _fetch_table(table_name: str) -> pd.DataFrame:
    if table_name not in _ALLOWED_TABLES:
        raise ValueError(f"Invalid table name: {table_name}")
    conn = connect_to_database()
    cur = conn.cursor(dictionary=True)
    try:
        cur.execute(f"SELECT * FROM `{table_name}`")
        rows = cur.fetchall()
        return pd.DataFrame(rows)
    finally:
        cur.close()
        conn.close()

# ----------------- Local Files Helpers ----------------- #
FILES_DIR = Path(__file__).resolve().parent / "files"


@st.cache_data(show_spinner=False)
def _discover_local_reports(folder: Path) -> dict[str, dict[int, Path]]:
    """
    Returns mapping:
      report_type -> {yy -> file_path}
    where yy is the 2-digit year suffix found at the end of the filename (" ... 26.csv").
    """
    out: dict[str, dict[int, Path]] = {}
    if not folder.exists():
        return out

    for p in folder.glob("*.csv"):
        name = p.name.lower()
        if name.startswith("stock summary report"):
            rtype = "stock_summary"
        elif name.startswith("sales by item"):
            rtype = "sales_by_item"
        elif name.startswith("purchases by item"):
            rtype = "purchases_by_item"
        else:
            continue

        m = re.search(r"\s(\d{2})\.csv$", p.name)
        if not m:
            continue
        yy = int(m.group(1))
        out.setdefault(rtype, {})[yy] = p

    return out


@st.cache_data(show_spinner=False)
def _read_local_csv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, encoding="utf-8-sig")


@st.cache_data(show_spinner=False)
def _db_available() -> tuple[bool, str]:
    try:
        conn = connect_to_database()
        conn.close()
        return True, ""
    except Exception as e:
        return False, str(e)


def _safe_replace_year(d: date, year: int) -> date:
    try:
        return d.replace(year=year)
    except ValueError:
        # Handle Feb 29 -> Feb 28
        return d.replace(year=year, day=28)

# ----------------- User Auth Helpers ----------------- #
def create_users_table():
    conn = connect_to_database()
    cur = conn.cursor()
    cur.execute("""
        CREATE TABLE IF NOT EXISTS users (
          username VARCHAR(50) PRIMARY KEY,
          password_hash CHAR(64) NOT NULL
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """)
    conn.commit()
    cur.close()
    conn.close()

def hash_password(pw: str) -> str:
    return hashlib.sha256(pw.encode()).hexdigest()

def register_user(username: str, password: str) -> bool:
    conn = connect_to_database()
    cur = conn.cursor()
    try:
        cur.execute(
            "INSERT INTO users(username, password_hash) VALUES (%s, %s)",
            (username, hash_password(password))
        )
        conn.commit()
        return True
    except mysql.connector.IntegrityError:
        return False
    finally:
        cur.close()
        conn.close()

def login_user(username: str, password: str) -> bool:
    conn = connect_to_database()
    cur = conn.cursor()
    cur.execute("SELECT password_hash FROM users WHERE username=%s", (username,))
    row = cur.fetchone()
    cur.close()
    conn.close()
    return bool(row and row[0] == hash_password(password))

_DB_OK, _DB_ERR = _db_available()
if _DB_OK:
    create_users_table()

# ----------------- Helpers for Downloads ----------------- #
def df_to_csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False).encode('utf-8')

def df_to_excel_bytes(df: pd.DataFrame) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Sheet1')
    return output.getvalue()

def df_to_pdf_bytes(df: pd.DataFrame) -> bytes:
    pdf = FPDF(orientation='L', unit='mm', format='A4')
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", size=7)

    padding = 2
    col_widths = []
    for col in df.columns:
        max_w = pdf.get_string_width(str(col)) + padding
        # Avoid O(n*m) scanning on large exports
        for val in df[col].head(200):
            w = pdf.get_string_width(str(val)) + padding
            if w > max_w:
                max_w = w
        col_widths.append(max_w)

    table_w = sum(col_widths)
    epw = pdf.w - 2 * pdf.l_margin
    if table_w > epw:
        scale = epw / table_w
        col_widths = [w * scale for w in col_widths]

    row_h = pdf.font_size * 1.2

    for i, header in enumerate(df.columns):
        pdf.cell(col_widths[i], row_h, str(header), border=1, align='C')
    pdf.ln(row_h)

    for _, row in df.iterrows():
        for i, cell in enumerate(row):
            pdf.cell(col_widths[i], row_h, str(cell), border=1)
        pdf.ln(row_h)

    out = pdf.output(dest='S')
    if isinstance(out, str):
        return out.encode('latin-1')
    return out

# ----------------- Layout & Title ----------------- #
st.set_page_config(page_title="General Report System", layout="wide")
st.markdown("<h1 style='text-align:center;'>📦 General Report & P&S Forecast</h1>", unsafe_allow_html=True)

# ----------------- Session State Init ----------------- #
for key, default in [
    ("logged_in", False),
    ("username", ""),
    ("report_df", pd.DataFrame()),
    ("forecast_df", pd.DataFrame()),
    ("use_db", False),
]:
    if key not in st.session_state:
        st.session_state[key] = default

# ----------------- Data Source ----------------- #
with st.sidebar:
    st.header("Data Source")
    st.session_state.use_db = st.toggle(
        "Use MySQL database",
        value=st.session_state.use_db if _DB_OK else False,
        help="Turn on only if you have the MySQL DB configured and populated. Otherwise use the local `files/` folder.",
        disabled=not _DB_OK,
    )
    if not _DB_OK:
        st.caption(f"DB unavailable: {_DB_ERR}")

# Local-files mode: bypass DB auth entirely.
if not st.session_state.use_db and not st.session_state.logged_in:
    st.session_state.logged_in = True
    st.session_state.username = "local"

# ----------------- Authentication ----------------- #
if not st.session_state.logged_in:
    c1, c2, c3 = st.columns([1,2,1])
    with c2:
        st.subheader("🔐 Please log in or register")
        mode = st.radio("", ["Login","Register"], horizontal=True)
        user = st.text_input("Username")
        pwd  = st.text_input("Password", type="password")
        if mode == "Register" and st.button("Create Account"):
            if register_user(user, pwd):
                st.success("Account created! Please log in.")
            else:
                st.error("Username already exists.")
        if mode == "Login" and st.button("Log In"):
            if login_user(user, pwd):
                st.session_state.logged_in = True
                st.session_state.username = user
                st.rerun()
            else:
                st.error("Invalid credentials.")
    st.stop()

# ----------------- Sidebar: CSV Import & Delete ----------------- #
with st.sidebar:
    st.write(f"👤 Logged in as **{st.session_state.username}**")
    if st.session_state.use_db and st.button("Log Out"):
        st.session_state.logged_in = False
        st.rerun()

    if st.session_state.use_db and st.session_state.username.lower() == "marcelin":
        st.header("📄 Import CSV to Table")
        base_table = st.selectbox("Table", ["stock_summary","sales_by_item","purchases_by_item"])
        year_opt   = st.selectbox("Report Year", ["Current Year","Previous Year"])
        replace_existing = st.checkbox("Replace existing table data before upload", value=False)
        uploaded   = st.file_uploader("Upload CSV", type=["csv"])
        if st.button("Upload to DB"):
            if uploaded:
                try:
                    df = pd.read_csv(uploaded)
                    # Replace NaN with None for proper SQL NULL handling
                    df = df.replace({np.nan: None})
                    if not df.empty:
                        table_name = base_table + ("" if year_opt=="Current Year" else "_previous")
                        conn = connect_to_database()
                        cur = conn.cursor()
                        try:
                            if replace_existing:
                                cur.execute(f"DELETE FROM `{table_name}`")

                            columns = list(df.columns)
                            cols_sql = ",".join(f"`{c}`" for c in columns)
                            placeholders = ",".join(["%s"] * len(columns))
                            insert_sql = f"INSERT INTO `{table_name}` ({cols_sql}) VALUES ({placeholders})"

                            def row_to_tuple(row_vals):
                                out = []
                                for v in row_vals:
                                    out.append(None if (v is None or pd.isna(v)) else v)
                                return tuple(out)

                            data = [row_to_tuple(r) for r in df.itertuples(index=False, name=None)]
                            cur.executemany(insert_sql, data)
                            conn.commit()
                            st.success(f"Successfully uploaded {len(data)} rows to {table_name}.")
                            _fetch_table.clear()
                        finally:
                            cur.close()
                            conn.close()
                    else:
                        st.warning("Empty file.")
                except Exception as e:
                    st.error(f"Upload error: {e}")

        st.header("🗑️ Delete Options")
        if st.button("Delete All Data"):
            table_name = base_table + ("" if year_opt=="Current Year" else "_previous")
            conn = connect_to_database()
            cur  = conn.cursor()
            try:
                cur.execute(f"DELETE FROM `{table_name}`")
                conn.commit()
                _fetch_table.clear()
            finally:
                cur.close()
                conn.close()
            st.success(f"Cleared {table_name}")

        if st.button("Clear Entire Database"):
            conn = connect_to_database()
            cur = conn.cursor()
            for t in [
                "stock_summary","sales_by_item","purchases_by_item",
                "stock_summary_previous","sales_by_item_previous","purchases_by_item_previous"
            ]:
                cur.execute(f"DELETE FROM `{t}`")
            conn.commit()
            cur.close()
            conn.close()
            _fetch_table.clear()
            st.success("All tables cleared!")
    else:
        if st.session_state.use_db:
            st.info("Welcome!") 
        else:
            st.info("Using local reports from `files/`.")

# ----------------- General Report Builder ----------------- #
def generate_general_report(
    *,
    stock: Optional[pd.DataFrame] = None,
    purchases: Optional[pd.DataFrame] = None,
    sales_cy: Optional[pd.DataFrame] = None,
    sales_py: Optional[pd.DataFrame] = None,
) -> pd.DataFrame:
    try:
        if stock is None or purchases is None or sales_cy is None or sales_py is None:
            stock = _fetch_table("stock_summary")
            purchases = _fetch_table("purchases_by_item")
            sales_cy = _fetch_table("sales_by_item")
            sales_py = _fetch_table("sales_by_item_previous")

        if stock.empty:
            st.warning("No data found in `stock_summary`. Upload your Zoho export first.")
            return pd.DataFrame()

        # ----------- Stock Summary -----------
        stock_sku_col = _find_col(stock, ["sku"])
        stock_id_col = _find_col(stock, ["item", "id"]) or _find_col(stock, ["item_id"])
        stock_item_col = _find_col(stock, ["item", "name"]) or _find_col(stock, ["item_name"])
        stock_cat_col = _find_col(stock, ["category", "name"]) or _find_col(stock, ["category"])
        opening_col = _find_col(stock, ["opening", "stock"]) or _find_col(stock, ["opening"])
        closing_stock_col = _find_col(stock, ["closing", "stock"]) or _find_col(stock, ["closing"])

        missing = [n for n, c in [
            ("SKU", stock_sku_col),
            ("Item Name", stock_item_col),
            ("Opening Stock", opening_col),
            ("Closing Stock", closing_stock_col),
        ] if c is None]
        if missing:
            st.error(
                "Stock Summary table is missing required columns: "
                + ", ".join(missing)
                + ". Check the CSV headers in Zoho export and re-upload."
            )
            return pd.DataFrame()

        stock = stock.copy()
        stock["SKU"] = stock[stock_sku_col].apply(_normalize_sku)
        stock["SKU_prefix"] = stock["SKU"].apply(_sku_prefix)
        stock["Item ID"] = stock[stock_id_col].astype(str).str.strip() if stock_id_col is not None else ""
        stock["Item Name"] = stock[stock_item_col].fillna("").astype(str).str.strip()
        stock = stock[(stock["SKU"] != "") & (stock["Item Name"] != "") & (stock["Item Name"].str.lower() != "nan")]
        stock["opening_balance"] = _to_number(stock[opening_col])
        stock["closing_stock"] = _to_number(stock[closing_stock_col])
        if stock_cat_col is not None:
            stock["category"] = stock[stock_cat_col].fillna("").astype(str).replace("nan", "").str.strip()
        else:
            stock["category"] = ""

        # Detect and coerce all remaining numeric columns in stock
        stock_opening_amt_col = _find_col(stock, ["opening", "stock", "amount"]) or _find_col(stock, ["opening", "amount"])
        stock_qty_in_col = _find_col(stock, ["quantity", "in"], ["amount"])
        stock_qty_in_amt_col = _find_col(stock, ["quantity", "in", "amount"])
        stock_qty_out_col = _find_col(stock, ["quantity", "out"], ["amount"])
        stock_qty_out_amt_col = _find_col(stock, ["quantity", "out", "amount"])
        stock_closing_amt_col = _find_col(stock, ["closing", "stock", "amount"]) or _find_col(stock, ["closing", "amount"])

        _extra_stock_cols: dict[str, str] = {}  # mapped_name -> original_col
        for mapped, col in [
            ("opening_stock_amount", stock_opening_amt_col),
            ("quantity_in", stock_qty_in_col),
            ("quantity_in_amount", stock_qty_in_amt_col),
            ("quantity_out", stock_qty_out_col),
            ("quantity_out_amount", stock_qty_out_amt_col),
            ("closing_stock_amount", stock_closing_amt_col),
        ]:
            if col is not None:
                stock[mapped] = _to_number(stock[col])
                _extra_stock_cols[mapped] = col

        # Build Item ID -> SKU_prefix mapping for merging other tables
        _id_to_prefix = (
            stock[stock["Item ID"] != ""][["Item ID", "SKU_prefix"]]
            .drop_duplicates("Item ID")
        )

        # Group items sharing the same first 6 SKU digits — sum ALL numeric cols
        _stock_agg: dict = {
            "SKU": ("SKU", "first"),
            "Item_Name": ("Item Name", "first"),
            "category": ("category", "last"),
            "opening_balance": ("opening_balance", "sum"),
            "closing_stock": ("closing_stock", "sum"),
        }
        for mapped in _extra_stock_cols:
            _stock_agg[mapped] = (mapped, "sum")

        stock_grouped = (
            stock.groupby("SKU_prefix", as_index=False)
            .agg(**_stock_agg)
        )
        stock_grouped = stock_grouped.rename(columns={"Item_Name": "Item Name"})

        # ----------- Purchases -----------
        purchases_grouped = pd.DataFrame({"SKU_prefix": stock_grouped["SKU_prefix"].unique(), "purchase": 0, "purchase_amount": 0.0})
        if not purchases.empty:
            purchases_sku_col = _find_col(purchases, ["sku"])
            purchases_id_col = _find_col(purchases, ["item", "id"]) or _find_col(purchases, ["item_id"])
            purchase_qty_col = (
                _find_col(purchases, ["quantity", "purchased"])
                or _find_col(purchases, ["qty", "purchased"])
                or _find_col(purchases, ["quantity"])
                or _find_col(purchases, ["purchase"])
            )
            purchase_amt_col = _find_col(purchases, ["amount"])

            def _prepare_purchases(p: pd.DataFrame) -> pd.DataFrame:
                p["purchase"] = _to_number(p[purchase_qty_col])
                p["purchase_amount"] = _to_number(p[purchase_amt_col]) if purchase_amt_col is not None else 0
                return p.groupby("SKU_prefix", as_index=False).agg(
                    purchase=("purchase", "sum"),
                    purchase_amount=("purchase_amount", "sum"),
                )

            if purchases_sku_col is not None and purchase_qty_col is not None:
                purchases = purchases.copy()
                purchases["SKU"] = purchases[purchases_sku_col].apply(_normalize_sku)
                purchases["SKU_prefix"] = purchases["SKU"].apply(_sku_prefix)
                purchases = purchases[purchases["SKU"] != ""]
                purchases_grouped = _prepare_purchases(purchases)
            elif purchases_id_col is not None and purchase_qty_col is not None:
                purchases = purchases.copy()
                purchases["Item ID"] = purchases[purchases_id_col].astype(str).str.strip()
                purchases = purchases[purchases["Item ID"] != ""]
                # Map Item ID -> SKU_prefix via stock data
                purchases = purchases.merge(_id_to_prefix, on="Item ID", how="left")
                purchases["SKU_prefix"] = purchases["SKU_prefix"].fillna("")
                purchases = purchases[purchases["SKU_prefix"] != ""]
                purchases_grouped = _prepare_purchases(purchases)

        # ----------- Sales (Current & Previous) -----------
        def build_sales_grouped(df_in: pd.DataFrame, qty_out: str, amount_out: str, cat_out: str) -> tuple[pd.DataFrame, Optional[str]]:
            if df_in.empty:
                return pd.DataFrame(), None

            sku_col = _find_col(df_in, ["sku"])
            id_col = _find_col(df_in, ["item", "id"]) or _find_col(df_in, ["item_id"])
            item_col = _find_col(df_in, ["item", "name"]) or _find_col(df_in, ["item_name"])
            cat_col = _find_col(df_in, ["category", "name"]) or _find_col(df_in, ["category"])
            qty_col = (
                _find_col(df_in, ["quantity", "sold"])
                or _find_col(df_in, ["qty", "sold"])
                or _find_col(df_in, ["quantity_sold"])
                or _find_col(df_in, ["quantity"])
            )
            amount_col = _find_col(df_in, ["amount"]) or _find_col(df_in, ["value"])

            if item_col is None or qty_col is None:
                return pd.DataFrame(), None

            out = df_in.copy()
            if id_col is not None:
                out["Item ID"] = out[id_col].astype(str).str.strip()
            out["Item Name"] = out[item_col].fillna("").astype(str).str.strip()
            out = out[(out["Item Name"] != "") & (out["Item Name"].str.lower() != "nan")]
            out[qty_out] = _to_number(out[qty_col])
            out[amount_out] = _to_number(out[amount_col]) if amount_col is not None else 0
            out[cat_out] = out[cat_col].fillna("").astype(str).replace("nan", "").str.strip() if cat_col is not None else ""

            # Prefer SKU-based grouping by first 6 digits
            if sku_col is not None:
                out["SKU"] = out[sku_col].apply(_normalize_sku)
                out["SKU_prefix"] = out["SKU"].apply(_sku_prefix)
                out = out[out["SKU"] != ""]
                grouped = (
                    out.groupby("SKU_prefix", as_index=False)
                    .agg(
                        **{
                            "Item Name": ("Item Name", "last"),
                            qty_out: (qty_out, "sum"),
                            amount_out: (amount_out, "sum"),
                            cat_out: (cat_out, "last"),
                        }
                    )
                )
                return grouped, "SKU_prefix"

            # Fall back: map Item ID -> SKU_prefix via stock data
            if id_col is not None and "Item ID" in out.columns:
                out = out[out["Item ID"] != ""]
                if not out.empty:
                    out = out.merge(_id_to_prefix, on="Item ID", how="left")
                    out["SKU_prefix"] = out["SKU_prefix"].fillna("")
                    out = out[out["SKU_prefix"] != ""]
                    if not out.empty:
                        grouped = (
                            out.groupby("SKU_prefix", as_index=False)
                            .agg(
                                **{
                                    "Item Name": ("Item Name", "last"),
                                    qty_out: (qty_out, "sum"),
                                    amount_out: (amount_out, "sum"),
                                    cat_out: (cat_out, "last"),
                                }
                            )
                        )
                        return grouped, "SKU_prefix"

            grouped = (
                out.groupby(["Item Name"], as_index=False)
                .agg(
                    **{
                        qty_out: (qty_out, "sum"),
                        amount_out: (amount_out, "sum"),
                        cat_out: (cat_out, "last"),
                    }
                )
            )
            return grouped, "Item Name"

        sales_cy_grouped, sales_cy_key = build_sales_grouped(
            sales_cy,
            qty_out="sales_current_year",
            amount_out="sales_value_current_year",
            cat_out="category_sales_current",
        )
        sales_py_grouped, sales_py_key = build_sales_grouped(
            sales_py,
            qty_out="sales_previous_year",
            amount_out="sales_value_previous_year",
            cat_out="category_sales_previous",
        )

        # ----------- Merge (all on SKU_prefix) -----------
        df = stock_grouped.merge(purchases_grouped, on="SKU_prefix", how="left")
        df["purchase"] = _to_number(df["purchase"])

        if not sales_cy_grouped.empty and sales_cy_key is not None:
            # Always drop Item Name from sales side to avoid _x/_y suffixes
            cy_merge = sales_cy_grouped.drop(columns=["Item Name"], errors="ignore")
            merge_key = "SKU_prefix" if sales_cy_key == "SKU_prefix" else "Item Name"
            if merge_key == "SKU_prefix":
                df = df.merge(cy_merge, on="SKU_prefix", how="left")
            else:
                df = df.merge(cy_merge, on="Item Name", how="left")
        else:
            df["sales_current_year"] = 0
            df["sales_value_current_year"] = 0
            df["category_sales_current"] = ""

        if not sales_py_grouped.empty and sales_py_key is not None:
            py_merge = sales_py_grouped.drop(columns=["Item Name"], errors="ignore")
            merge_key = "SKU_prefix" if sales_py_key == "SKU_prefix" else "Item Name"
            if merge_key == "SKU_prefix":
                df = df.merge(py_merge, on="SKU_prefix", how="left")
            else:
                df = df.merge(py_merge, on="Item Name", how="left")
        else:
            df["sales_previous_year"] = 0
            df["sales_value_previous_year"] = 0
            df["category_sales_previous"] = ""

        # Ensure all numeric columns are filled (NaN → 0) after left joins
        _all_num_cols = [
            "purchase", "purchase_amount",
            "sales_current_year", "sales_value_current_year",
            "sales_previous_year", "sales_value_previous_year",
        ] + list(_extra_stock_cols.keys())
        for _num_col in _all_num_cols:
            if _num_col in df.columns:
                df[_num_col] = _to_number(df[_num_col])
        # Ensure category text columns exist and have no NaN
        for _cat_col in ["category_sales_current", "category_sales_previous"]:
            if _cat_col not in df.columns:
                df[_cat_col] = ""
            else:
                df[_cat_col] = df[_cat_col].fillna("").astype(str)

        # Prefer stock category; fall back to sales categories (current then previous)
        df["category"] = df["category"].replace("nan", "").fillna("").astype(str).str.strip()
        df["category"] = df["category"].where(df["category"] != "", df["category_sales_current"].fillna(""))
        df["category"] = df["category"].where(df["category"] != "", df["category_sales_previous"].fillna(""))

        expected_closing = (df["opening_balance"] + df["purchase"] - df["sales_current_year"]).round(0)
        # Closing Balance should be the actual stock on hand from Stock Summary.
        df["closing_balance"] = _to_number(df["closing_stock"]).round(0)
        df["discrepancies"] = (expected_closing - df["closing_balance"]).round(0)

        df = df.rename(columns={"Item Name": "item_name"})

        # Build the ordered column list — core + extras from stock & purchases
        _output_cols = [
            "SKU",
            "item_name",
            "category",
            "opening_balance",
        ]
        if "opening_stock_amount" in df.columns:
            _output_cols.append("opening_stock_amount")
        if "quantity_in" in df.columns:
            _output_cols.append("quantity_in")
        if "quantity_in_amount" in df.columns:
            _output_cols.append("quantity_in_amount")
        _output_cols.append("purchase")
        if "purchase_amount" in df.columns:
            _output_cols.append("purchase_amount")
        if "quantity_out" in df.columns:
            _output_cols.append("quantity_out")
        if "quantity_out_amount" in df.columns:
            _output_cols.append("quantity_out_amount")
        _output_cols += [
            "sales_current_year",
            "sales_value_current_year",
            "sales_previous_year",
            "sales_value_previous_year",
            "closing_balance",
        ]
        if "closing_stock_amount" in df.columns:
            _output_cols.append("closing_stock_amount")
        _output_cols.append("discrepancies")

        df = df[_output_cols].copy()

        # Convert all numeric columns to int
        _int_cols = [c for c in _output_cols if c not in ("SKU", "item_name", "category")]
        for c in _int_cols:
            df[c] = _to_number(df[c]).round(0).astype(int)

        df = df.sort_values(["SKU", "item_name"]).reset_index(drop=True)
    except Exception as e:
        st.error(f"Error: {e}")
        df = pd.DataFrame()
    return df


def build_final_report(base_df: pd.DataFrame, as_of_date: date) -> pd.DataFrame:
    if base_df.empty:
        return base_df

    today = as_of_date
    year_start = date(today.year, 1, 1)
    total_weeks = 52
    days_elapsed = max((today - year_start).days + 1, 1)
    weeks_elapsed = min(max(days_elapsed / 7, 1), total_weeks)
    weeks_remaining = max(total_weeks - weeks_elapsed, 0)

    prev_as_of = _safe_replace_year(today, today.year - 1)
    prev_year_start = date(prev_as_of.year, 1, 1)
    prev_days_elapsed = max((prev_as_of - prev_year_start).days + 1, 1)
    prev_weeks_elapsed = min(max(prev_days_elapsed / 7, 1), total_weeks)

    _rename_map = {
        "purchase": "Purchases",
        "purchase_amount": "Purchase Amount",
        "sales_current_year": "Sales Current Year",
        "sales_value_current_year": "Sales Value Current Year",
        "sales_previous_year": "Previous Year Sales",
        "sales_value_previous_year": "Sales Value Previous year",
        "item_name": "Item Name",
        "category": "Category",
        "opening_balance": "Opening Balance",
        "opening_stock_amount": "Opening Stock Amount",
        "quantity_in": "Quantity In",
        "quantity_in_amount": "Quantity In Amount",
        "quantity_out": "Quantity Out",
        "quantity_out_amount": "Quantity Out Amount",
        "closing_balance": "Closing Balance",
        "closing_stock_amount": "Closing Stock Amount",
        "discrepancies": "Discrepancies",
    }
    f = base_df.rename(columns=_rename_map)

    # Build ordered column list — always include core, conditionally include extras
    _final_cols = ["SKU", "Item Name", "Category", "Opening Balance"]
    if "Opening Stock Amount" in f.columns:
        _final_cols.append("Opening Stock Amount")
    if "Quantity In" in f.columns:
        _final_cols.append("Quantity In")
    if "Quantity In Amount" in f.columns:
        _final_cols.append("Quantity In Amount")
    _final_cols.append("Purchases")
    if "Purchase Amount" in f.columns:
        _final_cols.append("Purchase Amount")
    if "Quantity Out" in f.columns:
        _final_cols.append("Quantity Out")
    if "Quantity Out Amount" in f.columns:
        _final_cols.append("Quantity Out Amount")
    _final_cols += [
        "Sales Current Year", "Sales Value Current Year",
        "Previous Year Sales", "Sales Value Previous year",
        "Closing Balance",
    ]
    if "Closing Stock Amount" in f.columns:
        _final_cols.append("Closing Stock Amount")
    _final_cols.append("Discrepancies")

    f = f[_final_cols].copy()

    current_weekly_rate = _to_number(f["Sales Current Year"]) / weeks_elapsed
    previous_weekly_rate = _to_number(f["Previous Year Sales"]) / prev_weeks_elapsed

    max_weekly_rate = np.maximum(current_weekly_rate, previous_weekly_rate)
    min_weekly_rate = np.minimum(current_weekly_rate, previous_weekly_rate)

    f["Max Forecast"] = (max_weekly_rate * weeks_remaining).clip(lower=0).round().astype(int)
    f["Min Forecast"] = (min_weekly_rate * weeks_remaining).clip(lower=0).round().astype(int)
    f["AVG Forecast"] = ((f["Max Forecast"] + f["Min Forecast"]) / 2).round().astype(int)

    f["Max Purchase Forecast"] = (f["Max Forecast"] - _to_number(f["Closing Balance"])).clip(lower=0).round().astype(int)
    f["Min Purchase Forecast"] = (f["Min Forecast"] - _to_number(f["Closing Balance"])).clip(lower=0).round().astype(int)
    f["AVG Purchase Forecast"] = ((f["Max Purchase Forecast"] + f["Min Purchase Forecast"]) / 2).round().astype(int)

    denom_stock = _to_number(f["Opening Balance"]) + _to_number(f["Purchases"])
    f["% Sales on Stock"] = np.where(denom_stock == 0, 0, _to_number(f["Sales Current Year"]) / denom_stock).round(4)

    den = _to_number(f["Previous Year Sales"])
    num = _to_number(f["Sales Current Year"])
    f["% Sales Difference"] = np.where(
        den == 0,
        np.where(num > 0, 1, 0),
        np.where(
            num == 0,
            np.where(den > 0, -1, 0),
            (num - den) / den
        )
    ).round(4)

    return f

# ----------------- Main ----------------- #
as_of_date = st.date_input(
    "Report as-of date (end date used in Zoho filters)",
    value=datetime.now().date(),
    help="Set this to the end date of the time-filtered Zoho exports you uploaded (weekly/YTD).",
    key="as_of_date",
)

local_reports = _discover_local_reports(FILES_DIR) if not st.session_state.use_db else {}
current_yy = None
previous_yy = None
if not st.session_state.use_db:
    with st.expander("Local `files/` inputs", expanded=False):
        if not local_reports:
            st.warning("No matching CSV reports found under `files/`.")
        else:
            all_years = sorted({y for m in local_reports.values() for y in m.keys()})
            if all_years:
                current_yy = st.selectbox("Current report year (YY)", all_years, index=len(all_years) - 1, key="current_yy")
                prev_years = [y for y in all_years if y < current_yy] or [current_yy]
                previous_yy = st.selectbox("Previous report year (YY)", prev_years, index=len(prev_years) - 1, key="previous_yy")
            for rt in ["stock_summary", "sales_by_item", "purchases_by_item"]:
                years = sorted(local_reports.get(rt, {}).keys())
                st.write(f"- {rt}: {', '.join(str(y) for y in years) if years else 'missing'}")


def _load_inputs_for_run() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    if st.session_state.use_db:
        return (
            _fetch_table("stock_summary"),
            _fetch_table("purchases_by_item"),
            _fetch_table("sales_by_item"),
            _fetch_table("sales_by_item_previous"),
        )

    if not local_reports:
        return pd.DataFrame(), pd.DataFrame(), pd.DataFrame(), pd.DataFrame()

    all_years = sorted({y for m in local_reports.values() for y in m.keys()})
    if not all_years:
        return pd.DataFrame(), pd.DataFrame(), pd.DataFrame(), pd.DataFrame()

    cy = int(st.session_state.get("current_yy", max(all_years)))
    py = int(st.session_state.get("previous_yy", max([y for y in all_years if y < cy], default=cy)))

    stock_p = local_reports.get("stock_summary", {}).get(cy)
    purch_p = local_reports.get("purchases_by_item", {}).get(cy)
    sales_cy_p = local_reports.get("sales_by_item", {}).get(cy)
    sales_py_p = local_reports.get("sales_by_item", {}).get(py)

    if not stock_p or not purch_p or not sales_cy_p or not sales_py_p:
        return pd.DataFrame(), pd.DataFrame(), pd.DataFrame(), pd.DataFrame()

    return (
        _read_local_csv(stock_p),
        _read_local_csv(purch_p),
        _read_local_csv(sales_cy_p),
        _read_local_csv(sales_py_p),
    )


if st.button("Generate Final Report"):
    stock_df, purch_df, sales_cy_df, sales_py_df = _load_inputs_for_run()
    if stock_df.empty or purch_df.empty or sales_cy_df.empty or sales_py_df.empty:
        if st.session_state.use_db:
            st.error("Missing required data in the database tables. Upload/populate the required tables and try again.")
        else:
            st.error("Missing required CSVs under `files/` for the selected years. Check the expander for which report types/years are available.")
        st.stop()
    st.session_state.report_df = generate_general_report(
        stock=stock_df,
        purchases=purch_df,
        sales_cy=sales_cy_df,
        sales_py=sales_py_df,
    )
    st.session_state.forecast_df = build_final_report(st.session_state.report_df, as_of_date)
    st.rerun()

if not st.session_state.report_df.empty:
    rpt = st.session_state.report_df.copy()
    st.subheader("📊 General Report")
    st.dataframe(rpt)

    fmt = st.selectbox("Download General Report As", ["CSV","Excel","PDF"], key="fmt_gen")
    if fmt == "CSV":
        st.download_button("📥 Download CSV", df_to_csv_bytes(rpt), "general_report.csv", "text/csv")
    elif fmt == "Excel":
        st.download_button("📥 Download XLSX", df_to_excel_bytes(rpt), "general_report.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        st.download_button("📥 Download PDF", df_to_pdf_bytes(rpt), "general_report.pdf", "application/pdf")

    st.caption("Change the as-of date above and click Generate Final Report to refresh the final report.")

if not st.session_state.forecast_df.empty:
    df = st.session_state.forecast_df.copy()
    st.subheader("📈 Final Report")
    st.dataframe(df)

    fmt2 = st.selectbox("Download Final Report As", ["CSV","Excel","PDF"], key="fmt_ps")
    if fmt2 == "CSV":
        st.download_button("📥 Download CSV", df_to_csv_bytes(df), "final_report.csv", "text/csv")
    elif fmt2 == "Excel":
        st.download_button("📥 Download XLSX", df_to_excel_bytes(df), "final_report.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        st.download_button("📥 Download PDF", df_to_pdf_bytes(df), "final_report.pdf", "application/pdf")

    cat_sales = df.groupby("Category")["Sales Current Year"].sum().reset_index()
    fig1 = px.pie(cat_sales, names="Category", values="Sales Current Year", title="Current Year Sales by Category")
    st.plotly_chart(fig1, use_container_width=True)

    cat_prev = df.groupby("Category")["Previous Year Sales"].sum().reset_index()
    fig2 = px.pie(cat_prev, names="Category", values="Previous Year Sales", title="Previous Year Sales by Category")
    st.plotly_chart(fig2, use_container_width=True)
